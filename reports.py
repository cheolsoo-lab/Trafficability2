# -*- coding: utf-8 -*-
"""
reports.py — PDF / Word 구조계산서 (엑셀 계산서와 동일 구성)
------------------------------------------------------------
1. 복토 두께에 따른 원지반에 작용하는 응력(σ) 산정      (3.1 형식)
2. 장비주행성 검토결과                                (3.1 형식)
3. Meyerhof and Hanna(1978) 허용지지력 검토           (2.1 형식)
4. 검토결과 그래프 및 결론

PDF 는 reportlab 으로 서버에서 직접 생성 → A4 폭 맞춤, 페이지 넘김 시 표 머리글 반복,
이미지/수식은 페이지 경계에서 잘리지 않음(KeepTogether).
"""

import io
import os

from PIL import Image as PILImage

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer,
                                Table, TableStyle)

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import core

H_CHUNK = 8   # 복토두께 열이 많으면 8개 단위로 표 분할


# ------------------------------------------------------------------
# 공통 데이터 가공
# ------------------------------------------------------------------
def f2(x):
    return f"{x:,.2f}"


def chunks(lst, n):
    return [lst[i:i + n] for i in range(0, len(lst), n)] or [[]]


def conclusion_lines(rows, p):
    ng = [r for r in rows if r["judge"] == "N.G"]
    lines = [f"검토 장비 {len(rows)}종 중 적합(O.K) {len(rows) - len(ng)}종, 부적합(N.G) {len(ng)}종으로 검토되었다."]
    if ng:
        lines.append("부적합 장비 : " + ", ".join(f"{r['name']}({r['spec']}, H={r['H']:.1f}m)" for r in ng)
                     + " → 복토두께 상향, 토목섬유 보강(T) 또는 장비 변경 검토가 필요하다.")
    else:
        lines.append("모든 장비가 제시된 복토두께에서 작용응력(σ) < 허용지지력(qa)를 만족하므로 장비 주행에 문제가 없는 것으로 판단된다.")
    lines.append(f"(적용 지반정수 : γ1={p['gamma1']:.1f} kN/m³, c2={p['c2']:.1f} kPa, φ1={p['phi1']:.1f}°, "
                 f"φ2={p['phi2']:.1f}°, θ={p['theta']:.1f}°, ε={p['eps']:.2f}, T={p['T']:.1f} kN/m, Fs={p['Fs']:.2f})")
    return lines


# ==================================================================
# PDF
# ==================================================================
_PDF_FONT = None


def _register_pdf_font():
    global _PDF_FONT
    if _PDF_FONT:
        return _PDF_FONT
    reg, bold = core.find_kr_font()
    try:
        if reg:
            kw = {"subfontIndex": 0} if reg.lower().endswith(".ttc") else {}
            pdfmetrics.registerFont(TTFont("KR", reg, **kw))
            if bold and bold != reg:
                kwb = {"subfontIndex": 0} if bold.lower().endswith(".ttc") else {}
                pdfmetrics.registerFont(TTFont("KR-B", bold, **kwb))
            else:
                pdfmetrics.registerFont(TTFont("KR-B", reg, **kw))
            _PDF_FONT = ("KR", "KR-B")
            return _PDF_FONT
    except Exception:
        pass
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    pdfmetrics.registerFont(UnicodeCIDFont("HYGothic-Medium"))
    _PDF_FONT = ("HYGothic-Medium", "HYGothic-Medium")
    return _PDF_FONT


def _img(data_or_path, max_w, max_h=None):
    if isinstance(data_or_path, (bytes, bytearray)):
        src = io.BytesIO(data_or_path)
        w, h = PILImage.open(io.BytesIO(data_or_path)).size
    else:
        src = data_or_path
        w, h = PILImage.open(data_or_path).size
    s = max_w / w
    if max_h and h * s > max_h:
        s = max_h / h
    return Image(src, width=w * s, height=h * s)


def build_pdf(project, rows, thicknesses, p):
    FN, FB = _register_pdf_font()
    buf = io.BytesIO()
    W, Hh = A4
    LM = RM = 15 * mm
    TW = W - LM - RM

    st = lambda name, **kw: ParagraphStyle(name, fontName=kw.pop("font", FN), **kw)  # noqa: E731
    sT = st("t", font=FB, fontSize=11.5, leading=15, spaceBefore=4, spaceAfter=4)
    sS = st("s", font=FB, fontSize=9.5, leading=13, spaceAfter=2)
    sN = st("n", fontSize=8.5, leading=12)
    sNs = st("ns", fontSize=7.5, leading=10, textColor=colors.HexColor("#444444"))
    sC = st("c", fontSize=7.5, leading=9, alignment=TA_CENTER)
    sCB = st("cb", font=FB, fontSize=7.5, leading=9, alignment=TA_CENTER)
    sL = st("l", fontSize=7.5, leading=9, alignment=TA_LEFT)

    def P(t, s=sC):
        return Paragraph(str(t).replace("\n", "<br/>"), s)

    def NM(t):
        return t if len(str(t)) <= 11 else Paragraph(str(t), sC)

    def heights(n_hdr_list, n_data, hdr_h=(13, 13)):
        return list(hdr_h[:len(n_hdr_list)]) + [10.5] * n_data

    base_style = [
        ("FONTNAME", (0, 0), (-1, -1), FN),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 0.8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0.8),
        ("LEFTPADDING", (0, 0), (-1, -1), 1.5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 1.5),
    ]

    def judge_styles(col, start, data_rows):
        out = []
        for i, r in enumerate(data_rows):
            c = colors.HexColor("#006100") if r["judge"] == "O.K" else colors.HexColor("#C00000")
            out.append(("TEXTCOLOR", (col, start + i), (col, start + i), c))
            out.append(("FONTNAME", (col, start + i), (col, start + i), FB))
            if r["judge"] == "N.G":
                out.append(("BACKGROUND", (col, start + i), (col, start + i), colors.HexColor("#FDE9E9")))
        return out

    def merged_name_styles(n_hdr, data_rows):
        """장비명 열(0)을 엑셀 병합처럼 보이게: 그룹 시작행에만 이름, 그룹 내부 가로선 없음"""
        sty = [("GRID", (1, 0), (-1, -1), 0.5, colors.black),
               ("BOX", (0, 0), (-1, -1), 0.8, colors.black),
               ("LINEAFTER", (0, 0), (0, -1), 0.5, colors.black),
               ("GRID", (0, 0), (0, n_hdr - 1), 0.5, colors.black)]
        names = [r["name"] for r in data_rows]
        for a, b, _ in core.group_spans(names):
            sty.append(("LINEABOVE", (0, n_hdr + a), (0, n_hdr + a), 0.5, colors.black))
        return sty

    story = []

    # ---------------- 1. 원지반 작용응력 ----------------
    story.append(Paragraph("1. 복토 두께에 따른 원지반에 작용하는 응력(σ)산정", sT))
    left = [Paragraph("장비 접지압에 대한 Model", st("h1", font=FB, fontSize=8.5, alignment=TA_CENTER)),
            _img(core.IMG_DISPERSION, TW * 0.47, 50 * mm)]
    legend = ("σ : 원지반상 작용응력(kPa)<br/>P : 장비접지압(kPa),  W : 장비중량(kN)<br/>H : 복토층의 두께(m)<br/>"
              "θ : 하중분산각(°)<br/>a, b : 접지길이, 접지폭(m)<br/>"
              f"ε : 충격계수,  γ1 : 복토 단위중량 = {p['gamma1']:.1f} kN/m³")
    right = [Paragraph("원지반상 작용응력(σ)", st("h2", font=FB, fontSize=8.5, alignment=TA_CENTER)),
             _img(core.formula_png("sigma"), TW * 0.47, 10.5 * mm),
             _img(core.formula_png("P"), TW * 0.38, 7.5 * mm),
             Spacer(1, 2 * mm), Paragraph(legend, st("lg", fontSize=8, leading=11.5))]
    t = Table([[left, right]], colWidths=[TW * 0.5, TW * 0.5])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(KeepTogether(t))
    story.append(Spacer(1, 3 * mm))

    fixed_w = [17, 10, 13, 10, 10, 15, 10, 10]   # mm
    for ci, hs in enumerate(chunks(list(thicknesses), H_CHUNK)):
        if ci > 0:
            story.append(Paragraph(f"(계속 — 복토 두께 {hs[0]:.1f} ~ {hs[-1]:.1f} m)", sNs))
        hw = (TW / mm - sum(fixed_w)) / max(len(hs), 1)
        colw = [w * mm for w in fixed_w] + [hw * mm] * len(hs)
        hdr1 = [P("장비명", sCB), P("규격", sCB), P("중량\n(W, kN)", sCB), P("장비", sCB), "", P("접지압\n(P, kPa)", sCB),
                P("하중\n분산각\n(θ)", sCB), P("충격\n계수\n(ε)", sCB),
                P("복토 두께에 따른 원지반상의 작용응력(σ, kN/m²)", sCB)] + [""] * (len(hs) - 1)
        hdr2 = ["", "", "", P("폭(b)", sCB), P("길이(a)", sCB), "", "", ""] + [P(f"{h:.1f}", sCB) for h in hs]
        data = [hdr1, hdr2]
        prev = None
        for r in rows:
            nm = r["name"] if r["name"] != prev else ""
            prev = r["name"]
            data.append([NM(nm), r["spec"], f2(r["W"]), f2(r["b"]), f"{r['a']:.3g}", f2(r["P"]),
                         f"{p['theta']:.1f}", f"{p['eps']:.2f}"] + [f2(r["sig_tbl"][h]) for h in hs])
        sty = base_style + merged_name_styles(2, rows) + [
            ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#F2F2F2")),
            ("SPAN", (0, 0), (0, 1)), ("SPAN", (1, 0), (1, 1)), ("SPAN", (2, 0), (2, 1)), ("SPAN", (3, 0), (4, 0)),
            ("SPAN", (5, 0), (5, 1)), ("SPAN", (6, 0), (6, 1)), ("SPAN", (7, 0), (7, 1)),
            ("SPAN", (8, 0), (8 + len(hs) - 1, 0)),
        ]
        tb = Table(data, colWidths=colw, repeatRows=2, rowHeights=[16, 13] + [10.5] * len(rows))
        tb.setStyle(TableStyle(sty))
        story.append(tb)
        story.append(Spacer(1, 2 * mm))

    # ---------------- 2. 장비주행성 검토결과 ----------------
    story.append(Spacer(1, 3 * mm))
    story.append(Paragraph("2. 장비주행성 검토결과", sT))
    colw = [w * mm for w in (24, 14, 24, 20, 26, 28, 18)]
    sc = TW / sum(colw)
    colw = [c * sc for c in colw]
    data = [[P("장비명", sCB), P("규격", sCB), P("접지압\n(P, kPa)", sCB), P("복토 두께\n(m)", sCB),
             P("작용응력\n(kPa)", sCB), P("허용지지력\n(kPa)", sCB), P("판정", sCB)]]
    prev = None
    for r in rows:
        nm = r["name"] if r["name"] != prev else ""
        prev = r["name"]
        data.append([NM(nm), r["spec"], f2(r["P"]), f"{r['H']:.1f}" + ("" if r["auto"] else " *"),
                     f2(r["sigma"]), f2(r["qa"]), r["judge"]])
    sty = base_style + merged_name_styles(1, rows) + [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F2F2"))]
    sty += judge_styles(6, 1, rows)
    tb = Table(data, colWidths=colw, repeatRows=1, rowHeights=[22] + [10.5] * len(rows))
    tb.setStyle(TableStyle(sty))
    story.append(tb)
    story.append(Spacer(1, 1.5 * mm))
    story.append(Paragraph("※ 허용지지력은 Meyerhof and Hanna(1978) 층상지반 지지력식에 의한 산정값 (3. 허용지지력 검토 참조). "
                           "복토 두께는 qa > σ 를 만족하는 최소 두께이며, * 표시는 검토자가 직접 지정한 두께임.", sNs))

    # ---------------- 3. Meyerhof and Hanna ----------------
    story.append(PageBreak())
    story.append(Paragraph("3. 원지반 허용지지력 검토", sT))
    story.append(Paragraph("▶ Meyerhof and Hanna(1978)모델", sS))
    story.append(Paragraph(": 층상지반에 대한 지지력 검토식으로 토목섬유의 인장력을 고려할 수 있도록 변형된 공식<br/>"
                           ": 연약한 점토층 위에 단단한 모래층이 있는 경우(c1=0, Φ2=0) 설계적용", sN))
    story.append(Spacer(1, 2 * mm))
    story.append(KeepTogether([_img(core.IMG_MEYERHOF, TW * 0.78, 56 * mm)]))
    story.append(Spacer(1, 2 * mm))
    eqs = [_img(core.formula_png("qa"), TW * 0.98, 12 * mm),
           Paragraph("* 원지반의 지지력 + 관입저항력 + 토목섬유 인장력", st("ri", fontSize=7.5, alignment=2)),
           Spacer(1, 1.5 * mm)]
    sub = Table([[_img(core.formula_png("Nq"), TW * 0.42, 8 * mm), _img(core.formula_png("Nc"), TW * 0.42, 8 * mm)],
                 [_img(core.formula_png("Fcs"), TW * 0.6, 8 * mm), ""]], colWidths=[TW * 0.5, TW * 0.5])
    sub.setStyle(TableStyle([("SPAN", (0, 1), (1, 1)), ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                             ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    eqs.append(sub)
    story.append(KeepTogether(eqs))
    story.append(Spacer(1, 2 * mm))

    Ks = core.ks_lookup(p["phi1"])
    lg = [
        ["여기서,", "a, b", ": 장비의 접지 길이, 접지폭", "", "", "B", ": 원지반상 하중 작용 폭", "", ""],
        ["", "γ1", ": 복토층 단위중량(kN/m³)", "=", f"{p['gamma1']:.2f}", "γ2", ": 원지반 단위중량(kN/m³)", "=", f"{p['gamma2']:.1f}"],
        ["", "Df", ": 근입심도(m)", "=", f"{p['Df']:.2f}", "c2", ": 원지반 점착력(kPa)", "=", f"{p['c2']:.1f}"],
        ["", "θ", ": 토목섬유가 수평면과 이루는 각", "=", f"{p['theta']:.1f}", "H", ": 복토층 두께", "", ""],
        ["", "Ks", ": 펀칭전단계수", "=", f"{Ks:.2f}", "Φ1", ": 복토층 내부마찰각", "=", f"{p['phi1']:.1f}"],
        ["", "i", ": 충격계수", "=", f"{p['eps']:.2f}", "Φ2", ": 원지반 내부마찰각", "=", f"{p['phi2']:.1f}"],
        ["", "T", ": 토목섬유의 허용인장력(kN/m)", "=", f"{p['T']:.1f}", "Fs", ": 안전율", "=", f"{p['Fs']:.2f}"],
        ["", "Nc(2)", f": {core.nc2(p['phi2']):.3f}", "", "", "Nq(2)", f": {core.nq2(p['phi2']):.3f}",
         "Nr(2)", f"{core.nr2(p['phi2']):.3f}"],
    ]
    t = Table(lg, colWidths=[w * mm for w in (13, 11, 50, 5, 13, 11, 47, 8, 13)])
    t.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), FN), ("FONTSIZE", (0, 0), (-1, -1), 8),
                           ("TOPPADDING", (0, 0), (-1, -1), 1), ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
                           ("ALIGN", (3, 0), (4, -1), "CENTER"), ("ALIGN", (7, 0), (8, -1), "CENTER")]))
    story.append(KeepTogether(t))
    story.append(Spacer(1, 2 * mm))
    kt = Table([["복토층 마찰각(Φ1˚)"] + [str(v) for v in core.KS_PHI], ["Ks"] + [f"{v:.2f}" for v in core.KS_VAL]],
               colWidths=[32 * mm] + [13 * mm] * 7)
    kt.setStyle(TableStyle(base_style + [("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                                         ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F2F2F2"))]))
    story.append(kt)
    story.append(Spacer(1, 3 * mm))

    story.append(Paragraph(f"■ 검토구간 (토목섬유 허용인장력 T = {p['T']:.1f} kN/m)", sS))
    colw = [w * mm for w in (30, 16, 12, 12, 17, 20, 18, 26, 14)]
    sc = TW / sum(colw)
    colw = [c * sc for c in colw]
    data = [[P("사용장비", sCB), P("장비 제원", sCB), "", "", P("원지반상 작용 응력", sCB), "",
             P("Geotextile\n허용강도(T)", sCB), P("허용 지지력(qa)", sCB), P("적합\n여부", sCB)],
            ["", P("접지압(P)", sCB), P("폭(b)", sCB), P("길이(L)", sCB), P("복토 두께(H)", sCB), P("작용 응력(σ)", sCB),
             "", P("Meyerhof and Hanna", sCB), ""]]
    for r in rows:
        data.append([NM(f"{r['name']}({r['spec']})"), f2(r["P"]), f2(r["b"]), f"{r['a']:.3g}", f"{r['H']:.1f}",
                     f2(r["sigma"]), f"{p['T']:.1f}", f2(r["qa"]), r["judge"]])
    sty = base_style + [("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                        ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#F2F2F2")),
                        ("SPAN", (0, 0), (0, 1)), ("SPAN", (1, 0), (3, 0)), ("SPAN", (4, 0), (5, 0)),
                        ("SPAN", (6, 0), (6, 1)), ("SPAN", (8, 0), (8, 1))]
    sty += judge_styles(8, 2, rows)
    tb = Table(data, colWidths=colw, repeatRows=2, rowHeights=[16, 13] + [10.5] * len(rows))
    tb.setStyle(TableStyle(sty))
    story.append(tb)

    # ---------------- 4. 그래프 / 결론 ----------------
    story.append(PageBreak())
    story.append(Paragraph("4. 검토결과 그래프 및 결론", sT))
    story.append(KeepTogether([_img(core.chart_png(rows), TW, 110 * mm)]))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph("▶ 검토 결론", sS))
    for ln in conclusion_lines(rows, p):
        story.append(Paragraph("· " + ln, sN))

    def on_page(cv, doc):
        cv.saveState()
        cv.setFont(FB, 9)
        cv.drawString(LM, Hh - 10 * mm, f"PROJECT : {project or ''}")
        cv.setFont(FN, 8)
        cv.drawRightString(W - RM, Hh - 10 * mm, "장비주행성 검토")
        cv.setLineWidth(0.6)
        cv.line(LM, Hh - 11.5 * mm, W - RM, Hh - 11.5 * mm)
        cv.drawCentredString(W / 2, 8 * mm, f"- {doc.page} -")
        cv.restoreState()

    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=LM, rightMargin=RM, topMargin=16 * mm,
                            bottomMargin=14 * mm, title="장비주행성 검토 구조계산서")
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


# ==================================================================
# Word (.docx)
# ==================================================================
def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _cell(cell, text, bold=False, size=8, color=None, align="center"):
    cell.text = ""
    par = cell.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == "center" else WD_ALIGN_PARAGRAPH.LEFT
    par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.space_before = Pt(0)
    run = par.add_run(str(text))
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _fix_layout(t, widths_cm):
    """표 고정 레이아웃 + 열너비 강제 + 좁은 셀 여백 (Word/한글/LibreOffice 공통으로 너비 유지)"""
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    mar = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), "40")
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)
    grid = t._tbl.tblGrid
    for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if j < len(widths_cm):
            gc.set(qn("w:w"), str(int(widths_cm[j] * 567)))
    for j, col in enumerate(t.columns):
        if j < len(widths_cm):
            col.width = Cm(widths_cm[j])
    for row in t.rows:
        for j, c in enumerate(row.cells):
            if j < len(widths_cm):
                c.width = Cm(widths_cm[j])


def _table(doc, data, widths_cm, hdr_rows=1, spans=(), judge_col=None, size=8):
    t = doc.add_table(rows=len(data), cols=len(data[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate(data):
        for j, v in enumerate(row):
            c = t.cell(i, j)
            c.width = Cm(widths_cm[j])
            is_h = i < hdr_rows
            color = None
            if judge_col is not None and j == judge_col and not is_h:
                color = "006100" if v == "O.K" else "C00000"
            _cell(c, v, bold=is_h or color is not None, size=size, color=color)
            if is_h:
                _shade(c, "F2F2F2")
            elif color == "C00000":
                _shade(c, "FDE9E9")
        if i < hdr_rows:
            _repeat_header(t.rows[i])
    _fix_layout(t, widths_cm)
    for (r1, c1, r2, c2) in spans:
        a = t.cell(r1, c1)
        txt = a.text
        m = a.merge(t.cell(r2, c2))
        _cell(m, txt, bold=r1 < hdr_rows, size=size)
        if r1 < hdr_rows:
            _shade(m, "F2F2F2")
    return t


def build_docx(project, rows, thicknesses, p):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.5)
    sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.4)
    TWc = 18.0

    sty = doc.styles["Normal"]
    sty.font.name = "맑은 고딕"
    sty.font.size = Pt(9)
    sty.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "맑은 고딕")

    hp = sec.header.paragraphs[0]
    hp.text = f"PROJECT : {project or ''}"
    hp.runs[0].bold = True
    hp.runs[0].font.size = Pt(9)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        r._r.append(el)

    def title(t, size=11.5):
        para = doc.add_paragraph()
        run = para.add_run(t)
        run.bold = True
        run.font.size = Pt(size)
        para.paragraph_format.space_before = Pt(6)
        para.paragraph_format.space_after = Pt(3)
        return para

    def text(t, size=8.5, italic=False, align=None):
        para = doc.add_paragraph()
        run = para.add_run(t)
        run.font.size = Pt(size)
        run.italic = italic
        para.paragraph_format.space_after = Pt(1)
        if align == "right":
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        return para

    def pic(data_or_path, width_cm, center=True):
        para = doc.add_paragraph()
        if center:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        src = io.BytesIO(data_or_path) if isinstance(data_or_path, (bytes, bytearray)) else data_or_path
        para.add_run().add_picture(src, width=Cm(width_cm))
        para.paragraph_format.keep_with_next = True
        return para

    # 1.
    title("1. 복토 두께에 따른 원지반에 작용하는 응력(σ)산정")
    lay = doc.add_table(rows=1, cols=2)
    lay.autofit = False
    lc, rc = lay.cell(0, 0), lay.cell(0, 1)
    _fix_layout(lay, [9.0, 9.0])
    _cell(lc, "장비 접지압에 대한 Model", bold=True, size=8.5)
    lc.add_paragraph().add_run().add_picture(core.IMG_DISPERSION, width=Cm(8.6))
    _cell(rc, "원지반상 작용응력(σ)", bold=True, size=8.5)
    rc.add_paragraph().add_run().add_picture(io.BytesIO(core.formula_png("sigma")), width=Cm(8.4))
    rc.add_paragraph().add_run().add_picture(io.BytesIO(core.formula_png("P")), width=Cm(6.0))
    for ln in ["σ : 원지반상 작용응력(kPa)", "P : 장비접지압(kPa),  W : 장비중량(kN)", "H : 복토층의 두께(m)",
               "θ : 하중분산각(°)", "a, b : 접지길이, 접지폭(m)",
               f"ε : 충격계수,  γ1 : 복토 단위중량 = {p['gamma1']:.1f} kN/m³"]:
        pr = rc.add_paragraph()
        pr.paragraph_format.space_after = Pt(0)
        pr.add_run(ln).font.size = Pt(8)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    fixed = [1.75, 0.95, 1.5, 1.0, 1.1, 1.55, 0.95, 0.95]
    for ci, hs in enumerate(chunks(list(thicknesses), H_CHUNK)):
        if ci > 0:
            text(f"(계속 — 복토 두께 {hs[0]:.1f} ~ {hs[-1]:.1f} m)", size=7.5, italic=True)
        hw = (TWc - sum(fixed)) / max(len(hs), 1)
        data = [["장비명", "규격", "중량\n(W, kN)", "장비", "", "접지압\n(P, kPa)", "하중\n분산각\n(θ)", "충격\n계수\n(ε)",
                 "복토 두께에 따른 원지반상의 작용응력(σ, kN/m²)"] + [""] * (len(hs) - 1),
                ["", "", "", "폭(b)", "길이(a)", "", "", ""] + [f"{h:.1f}" for h in hs]]
        for r in rows:
            data.append([r["name"], r["spec"], f2(r["W"]), f2(r["b"]), f"{r['a']:.3g}", f2(r["P"]),
                         f"{p['theta']:.1f}", f"{p['eps']:.2f}"] + [f2(r["sig_tbl"][h]) for h in hs])
        spans = [(0, 0, 1, 0), (0, 1, 1, 1), (0, 2, 1, 2), (0, 3, 0, 4), (0, 5, 1, 5), (0, 6, 1, 6), (0, 7, 1, 7)]
        if len(hs) > 1:
            spans.append((0, 8, 0, 8 + len(hs) - 1))
        for a, b, _ in core.group_spans([r["name"] for r in rows]):
            if b > a:
                spans.append((2 + a, 0, 2 + b, 0))
        _table(doc, data, fixed + [hw] * len(hs), hdr_rows=2, spans=spans, size=7)

    # 2.
    title("2. 장비주행성 검토결과")
    data = [["장비명", "규격", "접지압\n(P, kPa)", "복토 두께\n(m)", "작용응력\n(kPa)", "허용지지력\n(kPa)", "판정"]]
    for r in rows:
        data.append([r["name"], r["spec"], f2(r["P"]), f"{r['H']:.1f}" + ("" if r["auto"] else " *"),
                     f2(r["sigma"]), f2(r["qa"]), r["judge"]])
    spans = [(1 + a, 0, 1 + b, 0) for a, b, _ in core.group_spans([r["name"] for r in rows]) if b > a]
    _table(doc, data, [3.2, 1.8, 2.6, 2.3, 2.7, 3.0, 2.4], hdr_rows=1, spans=spans, judge_col=6, size=8)
    text("※ 허용지지력은 Meyerhof and Hanna(1978) 층상지반 지지력식에 의한 산정값 (3. 허용지지력 검토 참조). "
         "복토 두께는 qa > σ 를 만족하는 최소 두께이며, * 표시는 검토자가 직접 지정한 두께임.", size=7.5, italic=True)

    # 3.
    doc.add_page_break()
    title("3. 원지반 허용지지력 검토")
    title("▶ Meyerhof and Hanna(1978)모델", size=10)
    text(": 층상지반에 대한 지지력 검토식으로 토목섬유의 인장력을 고려할 수 있도록 변형된 공식")
    text(": 연약한 점토층 위에 단단한 모래층이 있는 경우(c1=0, Φ2=0) 설계적용")
    pic(core.IMG_MEYERHOF, 14.5)
    pic(core.formula_png("qa"), 17.5)
    text("* 원지반의 지지력 + 관입저항력 + 토목섬유 인장력", size=7.5, italic=True, align="right")
    sub = doc.add_table(rows=1, cols=3)
    for j, k in enumerate(["Nq", "Nc", "Fcs"]):
        c = sub.cell(0, j)
        c.width = Cm(6)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraphs[0].add_run().add_picture(io.BytesIO(core.formula_png(k)), width=Cm(5.6))
    Ks = core.ks_lookup(p["phi1"])
    lg = [
        ["여기서,", "γ1", ": 복토층 단위중량(kN/m³)", f"= {p['gamma1']:.2f}", "γ2", ": 원지반 단위중량(kN/m³)", f"= {p['gamma2']:.1f}"],
        ["", "Df", ": 근입심도(m)", f"= {p['Df']:.2f}", "c2", ": 원지반 점착력(kPa)", f"= {p['c2']:.1f}"],
        ["", "θ", ": 토목섬유가 수평면과 이루는 각", f"= {p['theta']:.1f}", "H", ": 복토층 두께", ""],
        ["", "Ks", ": 펀칭전단계수", f"= {Ks:.2f}", "Φ1", ": 복토층 내부마찰각", f"= {p['phi1']:.1f}"],
        ["", "i", ": 충격계수", f"= {p['eps']:.2f}", "Φ2", ": 원지반 내부마찰각", f"= {p['phi2']:.1f}"],
        ["", "T", ": 토목섬유의 허용인장력(kN/m)", f"= {p['T']:.1f}", "Fs", ": 안전율", f"= {p['Fs']:.2f}"],
        ["", "Nc(2)", f": {core.nc2(p['phi2']):.3f}", "", "Nq(2)", f": {core.nq2(p['phi2']):.3f}",
         f"Nr(2) : {core.nr2(p['phi2']):.3f}"],
    ]
    lt = doc.add_table(rows=len(lg), cols=7)
    for i, row in enumerate(lg):
        for j, v in enumerate(row):
            c = lt.cell(i, j)
            _cell(c, v, size=8, align="left")
    _fix_layout(lt, [1.4, 1.1, 5.4, 1.7, 1.1, 4.8, 2.5])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    _table(doc, [["복토층 마찰각(Φ1˚)"] + [str(v) for v in core.KS_PHI], ["Ks"] + [f"{v:.2f}" for v in core.KS_VAL]],
           [3.6] + [1.4] * 7, hdr_rows=0, size=8)
    title(f"■ 검토구간 (토목섬유 허용인장력 T = {p['T']:.1f} kN/m)", size=10)
    data = [["사용장비", "장비 제원", "", "", "원지반상 작용 응력", "", "Geotextile\n허용강도(T)", "허용 지지력(qa)", "적합\n여부"],
            ["", "접지압(P)", "폭(b)", "길이(L)", "복토 두께(H)", "작용 응력(σ)", "", "Meyerhof and Hanna", ""]]
    for r in rows:
        data.append([f"{r['name']}({r['spec']})", f2(r["P"]), f2(r["b"]), f"{r['a']:.3g}", f"{r['H']:.1f}",
                     f2(r["sigma"]), f"{p['T']:.1f}", f2(r["qa"]), r["judge"]])
    spans = [(0, 0, 1, 0), (0, 1, 0, 3), (0, 4, 0, 5), (0, 6, 1, 6), (0, 8, 1, 8)]
    _table(doc, data, [3.2, 1.8, 1.4, 1.4, 1.9, 2.2, 2.0, 2.6, 1.5], hdr_rows=2, spans=spans, judge_col=8, size=7.5)

    # 4.
    doc.add_page_break()
    title("4. 검토결과 그래프 및 결론")
    pic(core.chart_png(rows), 18)
    title("▶ 검토 결론", size=10)
    for ln in conclusion_lines(rows, p):
        text("· " + ln, size=9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
